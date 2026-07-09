from pypdf import PdfReader

from extensions import db
from models import UploadedDocument
from services.vocabulary import is_valid_vocab_word


def extract_text_from_file(file_storage):
    filename = file_storage.filename or "paste.txt"
    if filename.lower().endswith(".pdf"):
        reader = PdfReader(file_storage)
        pages = [page.extract_text() or "" for page in reader.pages]
        return filename, "\n".join(pages).strip()
    raw = file_storage.read()
    return filename, raw.decode("utf-8", errors="replace").strip()


def save_document(filename, text, language):
    doc = UploadedDocument(
        filename=filename,
        raw_text=text,
        language=language,
        word_count=len(text.split()),
    )
    db.session.add(doc)
    db.session.commit()
    return doc


def delete_document(doc):
    from models import AskHistory, DictionarySearch, VocabularyItem

    VocabularyItem.query.filter_by(document_id=doc.id).update({"document_id": None})
    DictionarySearch.query.filter_by(document_id=doc.id).update({"document_id": None})
    AskHistory.query.filter_by(document_id=doc.id).delete()
    db.session.delete(doc)
    db.session.commit()

import openpyxl
def keyword_search_chunks(text, query, context_chars=400):
    """Simple fallback until Phase 2b embeddings."""
    query_lower = query.lower()
    hits = []
    start = 0
    while True:
        idx = text.lower().find(query_lower, start)
        if idx == -1:
            break
        snippet = text[max(0, idx - context_chars): idx + context_chars]
        hits.append(snippet.strip())
        start = idx + len(query_lower)
    return hits[:5]


def parse_excel(file_storage):
    wb = openpyxl.load_workbook(file_storage, read_only=True)
    sheet = wb.active
    rows_iter = sheet.iter_rows(values_only=True)
    headers = [str(h).strip().lower() if h else "" for h in next(rows_iter)]
    rows = []
    for row in rows_iter:
        if not any(row):
            continue
        rows.append({headers[i]: (row[i] or "") for i in range(min(len(headers), len(row)))})
    return headers, rows


COLUMN_ALIASES = {
    "word": {"word", "term", "vocabulary", "front"},
    "meaning": {"meaning", "definition", "back", "translation"},
    "example": {"example", "sentence"},
    "topic": {"topic", "theme", "category"},
    "difficulty": {"difficulty", "level"},
    "notes": {"notes", "note"},
}


def guess_column_map(headers):
    mapping = {}
    normalized = {h: h.lower().strip() for h in headers}
    for field, aliases in COLUMN_ALIASES.items():
        for header, norm in normalized.items():
            if norm in aliases:
                mapping[field] = header
                break
    return mapping


def rows_to_cards(rows, column_map):
    cards = []
    for row in rows:
        word = str(row.get(column_map.get("word", ""), "")).strip()
        if not word or not is_valid_vocab_word(word):
            continue
        cards.append({
            "front": word,
            "back": str(row.get(column_map.get("meaning", ""), "")).strip(),
            "example": str(row.get(column_map.get("example", ""), "")).strip(),
            "topic": str(row.get(column_map.get("topic", ""), "")).strip(),
            "difficulty": str(row.get(column_map.get("difficulty", ""), "")).strip() or "beginner",
        })
    return cards